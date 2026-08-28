from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Iterable

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "manual-book"
SCREENSHOT_DIR = OUT_DIR / "screenshots"
SLICE_DIR = OUT_DIR / "generated-slices"
OUTPUT = OUT_DIR / "Manual Book Ecommerce Citra.docx"

BLUE = "2563EB"
LIGHT_BLUE = "EAF2FF"
NAVY = "172554"
GRAY = "64748B"
LIGHT_GRAY = "F1F5F9"
GREEN = "DCFCE7"
YELLOW = "FEF3C7"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None, size: int = 9) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: Iterable[Iterable[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        shade(table.rows[0].cells[index], BLUE)
        set_cell_text(table.rows[0].cells[index], header, True, "FFFFFF", 9)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if row_index % 2:
                shade(cells[index], "F8FAFC")
            set_cell_text(cells[index], str(value), False, NAVY, 8)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullets(doc: Document, items: Iterable[str], level: int = 0) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_note(doc: Document, title: str, text: str, fill: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell.margin_top = Cm(0.2)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{title}: ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    r = p.add_run(text)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_field_rules(doc: Document, rules: list[tuple[str, str, str]]) -> None:
    if not rules:
        return
    doc.add_heading("Data yang diisi dan aturannya", level=3)
    add_table(doc, ["Kolom", "Contoh isi", "Aturan sederhana"], rules, [1.35, 1.65, 3.45])


def image_slices(image_path: Path) -> list[Path]:
    SLICE_DIR.mkdir(parents=True, exist_ok=True)
    with Image.open(image_path) as im:
        width, height = im.size
        max_chunk_height = int(width * 1.08)
        if height <= max_chunk_height:
            return [image_path]
        outputs = []
        top = 0
        index = 1
        while top < height:
            bottom = min(top + max_chunk_height, height)
            crop = im.crop((0, top, width, bottom))
            target = SLICE_DIR / f"{image_path.stem}-bagian-{index}.png"
            crop.save(target, optimize=True)
            outputs.append(target)
            top = bottom
            index += 1
        return outputs


def add_screenshot(doc: Document, stem: str, caption: str) -> None:
    path = SCREENSHOT_DIR / f"{stem}.png"
    if not path.exists():
        add_note(doc, "Gambar belum tersedia", f"Screenshot {path.name} tidak ditemukan.", YELLOW)
        return
    parts = image_slices(path)
    for index, part in enumerate(parts, start=1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        p.add_run().add_picture(str(part), width=Inches(6.35))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(8)
        text = caption if len(parts) == 1 else f"{caption} — bagian {index} dari {len(parts)}"
        run = cap.add_run(text)
        run.italic = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GRAY)


def add_feature(
    doc: Document,
    title: str,
    purpose: str,
    steps: list[str],
    rules: list[tuple[str, str, str]],
    connections: list[str],
    screenshot: str | None = None,
    notes: list[str] | None = None,
) -> None:
    doc.add_heading(title, level=2)
    doc.add_paragraph(purpose)
    if screenshot:
        add_screenshot(doc, screenshot, f"Tampilan {title}")
    if steps:
        doc.add_heading("Cara pakai", level=3)
        add_numbered(doc, steps)
    add_field_rules(doc, rules)
    if connections:
        doc.add_heading("Data ini tersambung ke mana?", level=3)
        add_bullets(doc, connections)
    for note in notes or []:
        add_note(doc, "Perlu diperhatikan", note, YELLOW)


def add_toc(doc: Document) -> None:
    p = doc.add_paragraph()
    run = p.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, fld_sep, fld_end])
    doc.add_paragraph("Jika daftar isi belum muncul lengkap, klik kanan daftar isi di Microsoft Word lalu pilih Update Field.")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].font.color.rgb = RGBColor.from_string(NAVY)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for level, size in [(1, 22), (2, 16), (3, 12)]:
        style = styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE if level > 1 else NAVY)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
    header = section.header.paragraphs[0]
    header.text = "ECOMMERCE CITRA  •  MANUAL BOOK"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    add_page_number(section.footer.paragraphs[0])
    settings = doc.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def build_manual() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    run = p.add_run("MANUAL BOOK")
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Ecommerce Citra")
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Panduan pelanggan, admin, penjualan B2B, stok, konten, dan pengaturan").font.size = Pt(13)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    p.add_run("Versi dokumentasi: 28 Agustus 2026\nFormat: Microsoft Word, dapat diedit").font.color.rgb = RGBColor.from_string(GRAY)
    add_note(doc, "Tentang dokumen ini", "Panduan dibuat dari aplikasi yang berjalan, route, form, aturan validasi, model data, dan pengujian E2E. Screenshot memakai data contoh terpisah, jadi tidak mengubah data toko utama.")
    doc.add_page_break()

    doc.add_heading("Daftar Isi", level=1)
    add_toc(doc)
    doc.add_page_break()

    doc.add_heading("1. Cara Membaca Panduan Ini", level=1)
    doc.add_paragraph("Panduan ini memakai bahasa sehari-hari. Anda tidak perlu memahami istilah pemrograman. Cari nama menu yang ingin dipakai, ikuti langkahnya, lalu perhatikan tabel kolom dan aturan.")
    add_table(doc, ["Tanda", "Artinya"], [
        ("Wajib", "Harus diisi sebelum data bisa disimpan."),
        ("Opsional", "Boleh dikosongkan. Isi bila memang dibutuhkan."),
        ("Wajib jika...", "Baru wajib ketika pilihan tertentu diaktifkan, misalnya meminta faktur pajak."),
        ("Data snapshot", "Salinan data saat transaksi dibuat. Perubahan data master setelahnya tidak mengubah dokumen lama."),
    ], [1.4, 5.0])
    add_note(doc, "Akun contoh", "Screenshot memakai admin@citra.com dan akun pelanggan contoh. Password contoh tidak perlu dicantumkan pada panduan yang dibagikan ke pengguna akhir.", YELLOW)

    doc.add_heading("2. Peta Besar Sistem", level=1)
    doc.add_paragraph("Sistem memiliki dua jalur penjualan. Jalur B2C dipakai untuk belanja online biasa. Jalur B2B dipakai untuk penawaran dan dokumen penjualan perusahaan.")
    add_table(doc, ["Jalur", "Urutan sederhananya", "Hasil utama"], [
        ("B2C / toko online", "Produk → Keranjang/Beli Langsung → Checkout → Pembayaran → Diproses → Dikirim → Selesai", "Transaksi, invoice belanja, stok, poin, ulasan, retur, faktur pajak"),
        ("B2B dengan penawaran", "Quotation → Sales Order → Proforma (opsional) → Surat Jalan → Packing List → Invoice → Pembayaran", "Dokumen bisnis dengan nomor dan status masing-masing"),
        ("B2B order langsung", "Sales Order langsung → proses berikutnya sama seperti B2B di atas", "Dipakai saat harga sudah disepakati lewat chat/telepon"),
    ], [1.25, 2.65, 2.5])
    doc.add_heading("Hubungan data yang perlu dipahami", level=2)
    add_table(doc, ["Data sumber", "Dipakai oleh", "Dampak kalau diubah"], [
        ("Perusahaan aktif", "Produk, promo, transaksi, lokasi toko, dokumen B2B", "Admin hanya melihat data perusahaan yang sedang aktif."),
        ("Kategori + subkategori", "Produk dan filter katalog", "Produk tampil di kelompok yang sesuai."),
        ("Produk + varian", "Katalog, transaksi, stok, quotation, sales order", "Harga/stok master dipakai sebagai referensi; dokumen lama tetap menyimpan salinannya."),
        ("Lokasi toko", "Hitung ongkir", "Menjadi titik asal pengiriman."),
        ("Alamat pelanggan", "Checkout dan pengiriman", "Alamat terpilih disalin ke transaksi."),
        ("Transaksi", "Pembayaran, pengiriman, invoice, retur, ulasan, faktur pajak", "Status transaksi menentukan tombol tindakan yang tersedia."),
        ("Sales Order B2B", "Proforma, Surat Jalan, Packing List, Invoice B2B", "Menjadi sumber jumlah pesanan dan sisa yang belum dikirim/ditagih."),
        ("Role admin", "Menu dan tombol yang terlihat", "Akses dapat dibatasi per tindakan: lihat, buat, ubah, hapus, proses, kirim, dan lain-lain."),
    ], [1.4, 2.05, 2.95])

    doc.add_heading("3. Peran Pengguna", level=1)
    add_table(doc, ["Peran", "Biasanya mengerjakan", "Batas penting"], [
        ("Pengunjung", "Lihat katalog, cari produk, baca promo/blog, lacak order tamu", "Fitur akun seperti profil dan wishlist perlu login."),
        ("Pelanggan", "Belanja, checkout, alamat, pembayaran, poin, ulasan, retur, faktur pajak", "Hanya dapat melihat transaksi miliknya."),
        ("Admin/Super Admin", "Seluruh pengelolaan toko dan perusahaan", "Super admin dapat mengelola company, admin, dan role."),
        ("Sales", "Quotation dan Sales Order B2B", "Akses mengikuti role yang diberikan."),
        ("Gudang", "Stok, Surat Jalan, Packing List, proses pengiriman", "Tidak perlu diberi akses data sensitif yang tidak berkaitan."),
        ("Finance", "Pembayaran, invoice B2B, faktur pajak, laporan", "Nomor NPWP penuh sebaiknya hanya untuk role yang diizinkan."),
    ], [1.25, 2.7, 2.45])

    doc.add_heading("4. Panduan Pelanggan dan Pengunjung", level=1)
    add_feature(doc, "Beranda", "Halaman depan untuk melihat kategori populer, produk terbaru, rekomendasi, banner, promo, dan newsletter.",
                ["Gunakan kolom pencarian untuk mengetik nama produk/merek/kategori.", "Klik kategori atau kartu produk.", "Klik Detail untuk membuka informasi lengkap produk."], [],
                ["Banner berasal dari menu Admin > Banners.", "Produk dan stok berasal dari Admin > Products dan Stock.", "Bagian promo/flash sale mengikuti jadwal yang masih aktif."], "01-beranda")
    add_feature(doc, "Kategori dan pencarian produk", "Membantu pembeli menyaring produk berdasarkan kelompok dan spesifikasi.",
                ["Buka Semua Produk/Kategori.", "Pilih kategori atau centang filter yang sesuai.", "Gunakan urutan Terbaru/Terlaris bila dibutuhkan.", "Klik Terapkan Filter lalu buka Detail."], [],
                ["Filter teknis seperti diameter, panjang, grade, dan material berasal dari atribut varian produk.", "Produk tidak aktif tidak seharusnya muncul untuk pelanggan."], "02-kategori-produk")
    add_feature(doc, "Flash Sale", "Menampilkan produk yang mendapat harga khusus selama waktu tertentu.",
                ["Buka menu Flash Sale.", "Pastikan periode promo masih berjalan.", "Pilih produk dan lanjutkan ke detail/keranjang."], [],
                ["Harga promo, kuota, waktu mulai, dan waktu selesai diatur admin.", "Setelah waktu selesai atau kuota habis, harga promo tidak lagi dipakai."], "03-flash-sale")
    add_feature(doc, "Detail Produk", "Tempat memilih varian, melihat harga, stok, spesifikasi, dan ulasan sebelum membeli.",
                ["Pilih varian yang sesuai.", "Isi jumlah pembelian.", "Tambahkan catatan item jika perlu.", "Klik tambah ke keranjang atau beli langsung."], [
                    ("Varian", "M8 × 25 mm", "Wajib bila produk mempunyai lebih dari satu varian."),
                    ("Jumlah", "2", "Minimal 1 dan tidak boleh melebihi stok yang tersedia."),
                    ("Catatan", "Packing terpisah", "Opsional, maksimal 500 karakter pada proses checkout."),
                ], ["Pilihan varian mengacu ke SKU, harga, stok, berat, dan spesifikasi yang dibuat admin.", "Jumlah akan dipakai untuk menghitung subtotal dan berat ongkir."], "11-detail-produk")
    add_feature(doc, "Keranjang", "Menyimpan produk yang akan dibeli bersama-sama.",
                ["Login terlebih dahulu jika sistem mengarahkan ke halaman login.", "Ubah jumlah produk bila perlu.", "Hapus barang yang tidak jadi dibeli.", "Lanjutkan ke checkout."], [],
                ["Jumlah di keranjang belum memotong stok.", "Harga dan ketersediaan dicek kembali saat checkout."], "04-keranjang",
                ["Screenshot pengunjung menunjukkan pengalihan ke login. Ini memang aturan akses keranjang saat belum masuk."])
    add_feature(doc, "Checkout", "Mengumpulkan data penerima, pengiriman, kupon, pajak, dan metode pembayaran sebelum order dibuat.",
                ["Periksa barang dan jumlahnya.", "Pilih/tambahkan alamat penerima.", "Pilih layanan kirim.", "Masukkan kupon bila ada.", "Pilih metode pembayaran.", "Centang faktur pajak hanya bila diperlukan, lalu periksa total dan buat pesanan."], [
                    ("Nama penerima", "Budi Santoso", "Wajib, maksimal 100 karakter."),
                    ("Telepon", "+62 81234567890", "Wajib agar kurir dapat menghubungi penerima."),
                    ("Provinsi & kota", "Jawa Barat, Bandung", "Wajib; dipakai untuk ongkir."),
                    ("Alamat lengkap", "Jl. Melati No. 10...", "Wajib, maksimal 1.000 karakter."),
                    ("Kode pos", "40123", "Opsional, maksimal 12 karakter."),
                    ("Metode bayar", "BCA / BNI / BRI / Mandiri / CIMB / QRIS", "Wajib pilih satu."),
                    ("Kupon", "HEMAT10", "Opsional; harus aktif, masih berlaku, dan memenuhi minimum belanja/kuota."),
                    ("Nama/nomor/alamat NPWP", "PT Contoh / 00... / alamat", "Wajib hanya jika meminta faktur pajak."),
                ], ["Alamat tersimpan pada profil pelanggan dan disalin ke transaksi.", "Ongkir dihitung dari lokasi toko dan tujuan.", "Kupon mengurangi nilai sesuai jenis persen/nominal.", "PPN berasal dari pengaturan perusahaan.", "Checkout menghasilkan transaksi dan detail transaksi."], "16-checkout")
    add_feature(doc, "Login, daftar, dan lupa password", "Dipakai untuk masuk ke akun, membuat akun baru, atau meminta tautan penggantian password.",
                ["Untuk login, isi email dan password.", "Untuk daftar, isi nama, email, password, dan ulangi password.", "Untuk lupa password, isi email yang sudah terdaftar lalu buka tautan dari email."], [
                    ("Nama", "Budi Santoso", "Wajib, maksimal 255 karakter."),
                    ("Email", "budi@example.com", "Wajib, format email, tidak boleh sudah dipakai akun lain."),
                    ("Password daftar/reset", "minimal 6 karakter", "Wajib dan kolom konfirmasi harus sama."),
                    ("Ingat saya", "Dicentang", "Opsional; membuat sesi login bertahan lebih lama."),
                ], ["Akun pelanggan tersambung ke alamat, transaksi, poin, wishlist, notifikasi, retur, dan profil pajak."], "07-login")
    add_screenshot(doc, "08-registrasi", "Tampilan registrasi akun")
    add_screenshot(doc, "09-lupa-password", "Tampilan permintaan reset password")
    add_feature(doc, "Profil pelanggan dan alamat", "Tempat mengubah biodata, password, alamat, melihat poin, notifikasi, wishlist, dan riwayat pesanan.",
                ["Buka Profil.", "Pilih tab yang ingin dikelola.", "Simpan setiap bagian setelah diubah."], [
                    ("Nama depan", "Budi", "Wajib, maksimal 100 karakter."),
                    ("Email", "budi@example.com", "Wajib, valid, dan unik."),
                    ("Username", "budi.s", "Opsional, harus unik jika diisi."),
                    ("Tanggal lahir", "1990-01-01", "Opsional, harus berupa tanggal."),
                    ("Tautan sosial", "https://instagram.com/...", "Opsional, harus berupa URL."),
                    ("Bio", "Pemilik bengkel...", "Opsional, maksimal 1.000 karakter."),
                    ("Password baru", "minimal 8 karakter", "Pada menu profil: wajib konfirmasi dan password lama harus benar."),
                ], ["Perubahan profil tidak mengubah snapshot pelanggan pada transaksi/dokumen yang sudah dibuat.", "Alamat utama diprioritaskan saat checkout."], "12-profil-pelanggan")
    add_feature(doc, "Riwayat pesanan, notifikasi, wishlist, dan poin", "Kumpulan fitur setelah pelanggan login.",
                ["Buka tab Pesanan untuk melihat status dan detail.", "Buka Notifikasi untuk membaca perubahan order.", "Gunakan ikon hati untuk menyimpan produk.", "Buka Redeem Point untuk menukar poin dengan produk yang memenuhi syarat."], [],
                ["Pesanan berasal dari transaksi milik akun.", "Poin diberikan sesuai aturan setelah transaksi memenuhi status yang ditentukan.", "Produk redeem harus diaktifkan admin dan memiliki nilai poin."], "13-riwayat-pesanan")
    add_screenshot(doc, "14-wishlist", "Tampilan wishlist")
    add_screenshot(doc, "15-redeem-poin", "Tampilan penukaran poin")
    add_screenshot(doc, "17-notifikasi", "Tampilan notifikasi")
    add_feature(doc, "Lacak pesanan tanpa login", "Pengunjung/tamu dapat melihat order dengan email dan nomor order.",
                ["Buka Lacak Pesanan.", "Masukkan email yang dipakai saat checkout.", "Masukkan nomor order persis seperti pada email/halaman konfirmasi.", "Klik lacak."], [
                    ("Email", "tamu@example.com", "Wajib dan harus sama dengan email order."),
                    ("Nomor order", "ORD-...", "Wajib; sebaiknya salin-tempel agar tidak salah."),
                ], ["Hasil hanya menampilkan transaksi yang cocok dengan pasangan email + nomor order."], "05-lacak-pesanan",
                ["Percobaan dibatasi maksimal 3 kali per 15 menit. Pesan kesalahan dibuat umum agar tidak membocorkan data pelanggan."])
    add_feature(doc, "Retur, tukar barang, dan ulasan", "Diajukan dari detail transaksi setelah memenuhi syarat status order.",
                ["Buka detail transaksi.", "Pilih pengembalian dana atau tukar barang.", "Pilih item dan jumlah, tulis alasan, tambahkan foto bila perlu.", "Untuk ulasan, beri rating 1–5 dan komentar."], [
                    ("Jenis retur", "Refund / Exchange", "Wajib pilih satu."),
                    ("Alasan", "Barang rusak saat diterima", "Wajib, maksimal 1.000 karakter."),
                    ("Item & jumlah", "Baut M8 — 1", "Minimal satu item; jumlah tidak boleh melebihi yang dibeli."),
                    ("Foto retur", "foto-kemasan.jpg", "Opsional, maksimal 5 gambar; JPG/PNG/WebP; 4 MB per file."),
                    ("Rating", "5", "Wajib, angka 1 sampai 5."),
                    ("Foto ulasan", "hasil-pakai.jpg", "Opsional, maksimal 8 gambar; 4 MB per file."),
                ], ["Retur terkait langsung dengan transaksi dan baris item transaksi.", "Ulasan terkait produk yang benar-benar dibeli dan dapat dimoderasi admin."], None)
    add_feature(doc, "Blog dan halaman informasi", "Berisi artikel atau halaman seperti Tentang Kami, kebijakan, syarat, dan bantuan.",
                ["Buka Blog untuk daftar artikel.", "Klik judul artikel untuk membaca.", "Gunakan tautan footer untuk halaman informasi."], [],
                ["Isi dikelola melalui Admin > Konten Website."], "06-blog")
    add_feature(doc, "API katalog publik", "Dokumentasi untuk pihak yang perlu membaca data kategori dan produk lewat integrasi.",
                ["Buka dokumentasi API.", "Pilih endpoint yang diperlukan.", "Ikuti parameter dan contoh respons yang ditampilkan."], [],
                ["API membaca katalog yang sama dengan toko online dan hanya menyediakan endpoint yang tercantum."], "10-api-katalog-publik")

    doc.add_heading("5. Panduan Admin — Dasar", level=1)
    add_feature(doc, "Dashboard admin", "Ringkasan kondisi toko: penjualan, transaksi, pelanggan, stok, dan angka penting lainnya.",
                ["Pastikan perusahaan aktif di bagian atas sudah benar.", "Pilih periode bila tersedia.", "Klik kartu/ringkasan untuk masuk ke detail."], [],
                ["Angka dashboard mengikuti company aktif dan data transaksi pada periode terpilih."], "20-dashboard-admin")
    add_feature(doc, "Perusahaan aktif (multi-company)", "Memisahkan data operasional beberapa badan usaha/toko dalam satu aplikasi.",
                ["Buka Companies untuk membuat atau mengubah perusahaan.", "Gunakan pemilih company untuk berpindah ruang kerja.", "Periksa kembali company aktif sebelum memasukkan transaksi atau master data."], [
                    ("Nama perusahaan", "PT Citra Teknik Indonesia", "Wajib, maksimal 150 karakter."),
                    ("Nama legal", "PT Citra Teknik Indonesia", "Opsional."),
                    ("Prefix invoice", "CTI", "Wajib, unik, hanya huruf/angka/tanda dash atau underscore, maksimal 20."),
                    ("Email", "finance@contoh.co.id", "Opsional, harus format email."),
                    ("Logo", "logo.png", "Opsional, JPG/PNG/WebP, maksimal 2 MB."),
                ], ["Company aktif membatasi produk, promo, transaksi, banner, lokasi toko, dan dokumen B2B yang terlihat.", "Admin staff hanya dapat berpindah ke company yang ditugaskan kepadanya."], "52-perusahaan",
                ["Menonaktifkan company tidak sama dengan menghapus histori. Data lama tetap perlu dipertahankan."])
    add_feature(doc, "Admin Users dan Roles & Permissions", "Mengatur siapa yang bisa masuk admin serta menu/tindakan apa yang boleh dilakukan.",
                ["Buat role lebih dulu dan centang izin sesuai pekerjaan.", "Buat pengguna admin.", "Pilih tipe akun dan role.", "Jika perlu, atur role berbeda per company."], [
                    ("Nama admin", "Siti — Finance", "Wajib, maksimal 120 karakter."),
                    ("Email", "siti@contoh.co.id", "Wajib, valid, dan unik."),
                    ("Tipe akun", "Staff", "Wajib: super admin atau staff."),
                    ("Role", "Finance", "Wajib untuk staff agar punya izin kerja."),
                    ("Password", "minimal 8 karakter", "Wajib saat membuat akun; konfirmasi harus sama."),
                    ("Nama role", "Staff Gudang", "Wajib, unik, maksimal 100 karakter."),
                    ("Izin", "stock.index, delivery_notes.process", "Pilih hanya yang memang dibutuhkan."),
                ], ["Role menentukan menu dan tombol tindakan yang muncul.", "Penugasan company menentukan kumpulan data yang dapat dibuka staff."], "54-pengguna-admin")
    add_screenshot(doc, "55-role-izin", "Tampilan daftar role dan permission")
    add_feature(doc, "Ganti password admin", "Mengubah password akun admin yang sedang dipakai.",
                ["Isi password saat ini.", "Isi password baru minimal 8 karakter.", "Ulangi password baru dengan sama persis.", "Simpan."], [
                    ("Password saat ini", "••••••••", "Wajib dan harus cocok dengan akun."),
                    ("Password baru", "minimal 8 karakter", "Wajib."),
                    ("Konfirmasi", "sama dengan password baru", "Wajib sama persis."),
                ], [], "56-ganti-password")

    doc.add_heading("6. Panduan Admin — Master Produk dan Stok", level=1)
    add_feature(doc, "Kategori utama dan subkategori", "Menyusun produk agar mudah dicari pelanggan.",
                ["Buat kategori utama.", "Tambahkan subkategori dan pilih kategori induknya.", "Pilih kategori tersebut saat membuat produk."], [
                    ("Nama kategori utama", "Baut", "Wajib dan tidak boleh sama dengan kategori utama lain."),
                    ("Gambar kategori", "baut.png", "Opsional; JPG/PNG/WebP maksimal 4 MB atau URL gambar maksimal 2.048 karakter."),
                    ("Kategori induk", "Baut", "Wajib pada subkategori."),
                    ("Nama subkategori", "Baut Hex", "Wajib; harus unik di dalam induk yang sama."),
                ], ["Subkategori dipilih pada produk dan menjadi dasar menu/filter katalog."], "41-kategori-utama")
    add_screenshot(doc, "42-subkategori", "Tampilan subkategori")
    add_feature(doc, "Variants (daftar pilihan umum)", "Menyimpan pasangan nama dan nilai seperti Warna = Hitam atau Ukuran = M8.",
                ["Klik tambah varian.", "Isi nama kelompok dan nilainya.", "Simpan lalu gunakan pada produk bila dibutuhkan."], [
                    ("Nama", "Warna", "Wajib, maksimal 255 karakter."),
                    ("Nilai", "Hitam", "Wajib, maksimal 255 karakter; pasangan nama+nilai tidak boleh duplikat."),
                ], ["Varian umum membantu menjaga penulisan pilihan tetap konsisten."], "43-varian")
    add_feature(doc, "Products", "Master utama barang yang dijual, termasuk varian, SKU, harga, stok, berat, gambar, dan spesifikasi.",
                ["Klik Create Product.", "Isi nama dan deskripsi.", "Pilih/tambah kategori.", "Isi minimal satu varian.", "Periksa harga, stok, dan berat.", "Isi spesifikasi teknis agar filter katalog bekerja.", "Aktifkan status lalu simpan."], [
                    ("Nama produk", "Baut Hex M8 × 25 mm", "Wajib, maksimal 255 karakter."),
                    ("Kategori", "Baut > Baut Hex", "Opsional secara sistem, tetapi sangat disarankan agar produk mudah ditemukan."),
                    ("Status", "Active", "Wajib. Inactive menyembunyikan produk dari penjualan."),
                    ("Varian", "minimal 1", "Wajib minimal satu varian."),
                    ("SKU", "BH-M8-25-GV", "Terbentuk/diisi sebagai identitas varian; sebaiknya unik dan konsisten."),
                    ("Harga", "850", "Wajib, angka 0 atau lebih."),
                    ("Stok", "100", "Wajib, bilangan bulat 0 atau lebih."),
                    ("Berat", "100 gram", "Wajib, minimal 1 gram; dipakai untuk ongkir."),
                    ("P/L/T", "10 / 5 / 3 cm", "Opsional, angka 0 atau lebih untuk referensi packing."),
                    ("Gambar varian", "baut.webp", "Opsional; JPG/PNG/WebP, maksimal 4 MB."),
                    ("Redeem point", "500 poin", "Jika produk redeem diaktifkan, nilai poin wajib minimal 1."),
                ], ["Produk dipakai oleh katalog, keranjang, checkout, quotation, sales order, stok, promo, dan laporan.", "Dokumen transaksi menyimpan snapshot nama/harga sehingga perubahan master tidak mengubah transaksi lama."], "40-form-produk")
    add_screenshot(doc, "39-produk", "Tampilan daftar produk")
    add_feature(doc, "Import produk Excel", "Memasukkan banyak produk sekaligus menggunakan template resmi sistem.",
                ["Unduh template dari halaman Products.", "Jangan mengubah nama atau urutan kolom.", "Isi baris sesuai contoh.", "Unggah file .xlsx/.xls lalu periksa pesan hasil import."], [
                    ("File", "produk-baru.xlsx", "Wajib, hanya .xlsx atau .xls."),
                    ("price, stock, weight_grams", "850, 100, 100", "Wajib pada setiap baris."),
                    ("Atribut teknis", "diameter, length_mm, thread_type, grade, material", "Kolom template harus tersedia; isi konsisten agar filter bekerja."),
                ], ["Import membuat/memperbarui produk, varian, kategori terkait, dan atribut teknis sesuai isi template."], None,
                ["Selalu coba dengan beberapa baris dahulu sebelum mengimpor file besar."])
    add_feature(doc, "Stock", "Melihat stok per varian, menetapkan batas stok rendah, dan mencatat penyesuaian masuk/keluar.",
                ["Cari varian produk.", "Untuk koreksi, pilih Masuk atau Keluar.", "Isi jumlah dan alasan.", "Atur batas stok rendah bila perlu.", "Simpan dan periksa riwayat pergerakan."], [
                    ("Varian produk", "BH-M8-25-GV", "Wajib dan harus berasal dari company aktif."),
                    ("Tipe", "Masuk / Keluar", "Wajib."),
                    ("Jumlah", "25", "Wajib, bilangan bulat minimal 1."),
                    ("Keterangan", "Hasil stok opname", "Opsional, tetapi disarankan untuk jejak audit."),
                    ("Batas stok rendah", "10", "Wajib saat diubah, minimal 0."),
                ], ["Checkout/pengiriman B2B dapat mengurangi stok.", "Retur/penyesuaian dapat menambah stok sesuai proses.", "Stock movement menjadi sumber riwayat dan laporan stok."], "44-stok",
                ["Jangan memakai penyesuaian stok untuk menyamarkan selisih tanpa keterangan. Tuliskan alasan yang mudah diaudit."])

    doc.add_heading("7. Panduan Admin — Promo, Pelanggan, dan Konten", level=1)
    add_feature(doc, "Membership Tiers", "Mengelompokkan pelanggan berdasarkan total belanja dan memberi penjelasan manfaat.",
                ["Buat level dari yang paling rendah.", "Isi minimum spending dan warna.", "Urutkan level.", "Simpan."], [
                    ("Nama", "Gold", "Wajib, unik, maksimal 100 karakter."),
                    ("Minimum spending", "10000000", "Wajib, bilangan bulat minimal 0."),
                    ("Warna", "#F59E0B", "Wajib, maksimal 30 karakter."),
                    ("Manfaat", "Prioritas layanan", "Opsional."),
                    ("Urutan", "3", "Opsional, minimal 0."),
                ], ["Tier dapat diperbarui berdasarkan total belanja pelanggan dan terlihat pada data customer."], "38-form-level-member")
    add_screenshot(doc, "37-level-member", "Tampilan daftar level member")
    add_feature(doc, "Customers", "Daftar akun pelanggan beserta data kontak, poin, tier, dan aktivitasnya.",
                ["Cari pelanggan berdasarkan nama/email.", "Buka data yang diperlukan.", "Gunakan data pelanggan saat membuat transaksi manual atau dokumen B2B."], [],
                ["Customer dapat dipilih pada transaksi manual, quotation, dan sales order.", "Dokumen menyimpan snapshot data pelanggan agar histori tidak berubah."], "36-pelanggan")
    add_feature(doc, "Flash Sale admin", "Membuat promo waktu terbatas untuk varian tertentu.",
                ["Buat event flash sale.", "Isi waktu mulai dan selesai.", "Tambahkan minimal satu varian, harga promo, dan kuota.", "Aktifkan saat siap."], [
                    ("Nama event", "Flash Sale Akhir Bulan", "Wajib."),
                    ("Mulai", "2026-08-28 10:00", "Wajib."),
                    ("Selesai", "2026-08-28 22:00", "Wajib dan harus setelah waktu mulai."),
                    ("Status", "Draft / Active / Inactive", "Wajib."),
                    ("Varian", "BH-M8-25-GV", "Minimal satu; varian tidak boleh duplikat dalam satu event."),
                    ("Harga promo", "700", "Wajib, angka 0 atau lebih."),
                    ("Kuota", "50", "Wajib, bilangan bulat minimal 1."),
                ], ["Harga promo muncul di detail/katalog selama aktif dan dalam periode.", "Kuota membatasi jumlah yang bisa memakai promo."], "45-flash-sale-admin")
    add_feature(doc, "Coupons", "Memberi diskon berdasarkan kode saat checkout.",
                ["Buat kode dan nama kupon.", "Pilih diskon persen atau nominal tetap.", "Atur nilai, minimum belanja, kuota, periode, dan status.", "Uji kode pada checkout."], [
                    ("Kode", "HEMAT10", "Wajib, unik, maksimal 50 karakter."),
                    ("Nama", "Diskon Pelanggan Baru", "Wajib, maksimal 100 karakter."),
                    ("Jenis", "Percent / Fixed", "Wajib."),
                    ("Nilai", "10", "Wajib, bilangan bulat minimal 1."),
                    ("Maksimum diskon", "50000", "Opsional; berguna untuk diskon persen."),
                    ("Minimum belanja", "200000", "Opsional, minimal 0."),
                    ("Batas penggunaan", "100", "Opsional, minimal 1."),
                    ("Periode", "1–30 September", "Tanggal selesai tidak boleh sebelum tanggal mulai."),
                    ("Member only", "Ya", "Jika aktif, tamu/non-member tidak dapat memakai kupon."),
                ], ["Kupon divalidasi ulang saat checkout dan diskonnya disimpan pada transaksi."], "46-kupon")
    add_feature(doc, "Banners", "Mengatur gambar besar dan gambar samping pada beranda.",
                ["Pilih tipe carousel atau side.", "Unggah gambar atau isi URL gambar.", "Isi tautan tujuan bila banner bisa diklik.", "Atur urutan dan aktifkan."], [
                    ("Tipe", "Carousel / Side", "Wajib."),
                    ("Gambar", "banner.webp", "Gunakan file atau URL; JPG/PNG/WebP maksimal 12 MB sebelum kompresi."),
                    ("Target URL", "https://.../promo/...", "Opsional, harus URL lengkap."),
                    ("Urutan", "1", "Wajib, bilangan bulat minimal 1."),
                    ("Aktif", "Ya", "Hanya banner aktif yang ditampilkan."),
                ], ["Banner tampil pada beranda company aktif."], "48-banner")
    add_feature(doc, "Newsletter", "Mengelola pendaftar email, ekspor daftar, dan pengiriman kampanye.",
                ["Pantau subscriber aktif.", "Buat/kirim kampanye hanya ke penerima yang sesuai.", "Gunakan ekspor bila perlu diproses resmi di luar sistem."], [
                    ("Email subscriber", "customer@example.com", "Wajib dan harus format email."),
                    ("Subjek kampanye", "Promo September", "Wajib saat mengirim kampanye."),
                    ("Isi", "Ringkasan promo...", "Wajib; hindari data sensitif."),
                ], ["Form newsletter di beranda menambah subscriber.", "Unsubscribe menonaktifkan penerimaan kampanye berikutnya."], "49-newsletter")
    add_feature(doc, "Promo Pages", "Membuat landing page khusus kampanye.",
                ["Isi judul, subjudul, deskripsi, dan gambar.", "Tambahkan tombol ajakan dan URL tujuan.", "Atur periode serta status aktif.", "Simpan dan buka URL promo untuk mengecek hasil."], [
                    ("Judul", "Promo Proyek Akhir Tahun", "Wajib, maksimal 255 karakter."),
                    ("Slug", "promo-akhir-tahun", "Opsional; jika diisi harus unik."),
                    ("Hero image", "promo.webp", "Opsional; JPG/PNG/WebP maksimal 6 MB atau URL gambar."),
                    ("CTA label", "Belanja Sekarang", "Opsional, maksimal 50 karakter."),
                    ("CTA URL", "https://.../kategori", "Opsional, harus URL."),
                    ("Periode", "1–31 Desember", "Tanggal selesai tidak boleh sebelum mulai."),
                ], ["Halaman dapat dibuka lewat /promo/slug dan ditautkan dari banner/newsletter."], "50-halaman-promo")
    add_feature(doc, "Konten Website", "Membuat halaman informasi dan artikel blog.",
                ["Pilih tipe Page atau Post.", "Isi judul dan isi konten.", "Tambahkan ringkasan, gambar, dan info SEO bila perlu.", "Atur tanggal terbit dan aktifkan."], [
                    ("Tipe", "Page / Post", "Wajib."),
                    ("Judul", "Panduan Memilih Baut", "Wajib, maksimal 255 karakter."),
                    ("Slug", "panduan-memilih-baut", "Opsional; harus unik jika diisi."),
                    ("Ringkasan", "Cara singkat...", "Opsional, maksimal 500 karakter."),
                    ("Hero image", "artikel.webp", "Opsional; maksimal 6 MB atau URL valid."),
                    ("Meta description", "Panduan praktis...", "Opsional, maksimal 500 karakter."),
                ], ["Post tampil di Blog; Page dapat dibuka melalui slug halaman."], "51-konten-website")

    doc.add_heading("8. Panduan Admin — Pengaturan Toko dan Integrasi", level=1)
    add_feature(doc, "Store Settings", "Mengatur identitas toko, lokasi asal, rekening manual, pajak, WhatsApp, dan media sosial.",
                ["Pilih tab pengaturan.", "Isi satu bagian hingga lengkap.", "Simpan bagian tersebut sebelum berpindah tab.", "Uji perubahan di halaman pelanggan."], [
                    ("Nama toko", "Ecommerce Citra", "Wajib, maksimal 120 karakter."),
                    ("Logo", "logo.png", "Opsional; JPG/PNG/WebP maksimal 2 MB."),
                    ("Provinsi & kota asal", "DKI Jakarta — Jakarta Barat", "Wajib pada lokasi toko; dipakai untuk ongkir."),
                    ("Nama bank", "BCA", "Wajib jika transfer manual digunakan."),
                    ("Nomor rekening", "1234567890", "Wajib jika transfer manual digunakan."),
                    ("Nama pemilik rekening", "PT Citra...", "Wajib jika transfer manual digunakan."),
                    ("PPN", "11%", "Nilai wajib 0–100, maksimal 2 angka desimal; nama pajak wajib jika pajak aktif."),
                    ("Media sosial", "https://instagram.com/...", "Opsional, harus URL lengkap."),
                    ("WhatsApp store ID", "citra-store", "Wajib untuk gateway; hanya huruf, angka, titik, underscore, dash."),
                ], ["Nama/logo tampil pada pelanggan dan dokumen.", "Lokasi toko menjadi asal hitung ongkir.", "Rekening ditampilkan untuk transfer manual.", "PPN memengaruhi perhitungan checkout/dokumen baru.", "Pengaturan berlaku per company pada bagian yang memang company-scoped."], "47-pengaturan-toko",
                ["Perubahan tarif pajak tidak mengubah transaksi lama karena transaksi menyimpan nilai pajak saat dibuat."])
    add_feature(doc, "API Katalog admin", "Memberi panduan endpoint katalog untuk integrator.",
                ["Baca daftar endpoint dan parameter.", "Uji dari sistem integrasi dengan batas penggunaan yang wajar.", "Jangan menganggap API ini menyediakan fungsi admin atau checkout jika tidak tertulis."], [],
                ["Data berasal dari kategori, produk, dan varian aktif."], "53-api-katalog-admin")

    doc.add_heading("9. Panduan Admin — Transaksi B2C", level=1)
    add_feature(doc, "Buat Transaksi Manual", "Dipakai untuk order yang diterima lewat telepon, chat, atau counter tanpa checkout pelanggan.",
                ["Pilih customer existing atau manual.", "Tambahkan minimal satu produk/varian.", "Isi jumlah, harga, dan diskon per item bila ada.", "Isi diskon total, ongkir, PPN, pembayaran, pengiriman, dan faktur pajak bila dibutuhkan.", "Periksa grand total lalu simpan."], [
                    ("Mode customer", "Existing / Manual", "Wajib."),
                    ("Customer existing", "Budi Santoso", "Wajib jika mode Existing."),
                    ("Nama & telepon manual", "Toko Maju / 0812...", "Wajib jika mode Manual; email opsional tetapi harus valid."),
                    ("Item", "BH-M8-25-GV", "Minimal satu varian produk."),
                    ("Qty", "10", "Wajib, bilangan bulat minimal 1 dan tidak boleh melebihi stok saat diproses."),
                    ("Harga satuan", "850", "Wajib, angka 0 atau lebih."),
                    ("Diskon/ongkir", "5000 / 20000", "Opsional, angka 0 atau lebih."),
                    ("PPN", "11", "Opsional, 0–100."),
                    ("Data NPWP", "nama, nomor, alamat", "Wajib hanya jika opsi faktur pajak dicentang."),
                ], ["Menghasilkan Transaction ber-source manual dan Transaction Detail.", "Item menyimpan snapshot nama, SKU, dan harga.", "Status pembayaran manual dipisahkan dari status pesanan.", "Stok mengikuti tahap proses transaksi, bukan sekadar saat draft dibuat."], "29-buat-transaksi-manual")
    add_feature(doc, "Transactions", "Pusat pemrosesan semua order B2C dari checkout maupun input manual.",
                ["Gunakan pencarian dan filter status.", "Buka detail transaksi.", "Verifikasi pembayaran manual bila ada.", "Proses order setelah pembayaran sesuai aturan.", "Isi kurir/resi lalu tandai dikirim.", "Selesaikan atau batalkan sesuai kondisi."], [],
                ["Perubahan status dicatat pada histori transaksi dan dapat memicu notifikasi.", "Saat diproses, stok varian dikurangi dan stock movement dicatat.", "Data transaksi dipakai oleh invoice, label pengiriman, retur, ulasan, laporan, dan faktur pajak."], "30-transaksi",
                ["Jangan memproses order yang pembayaran manualnya belum diverifikasi, kecuali prosedur bisnis memang mengizinkan dan admin punya wewenang."])
    add_feature(doc, "Return Requests admin", "Menilai permintaan refund atau tukar barang dari pelanggan.",
                ["Buka permintaan dan cocokkan dengan transaksi.", "Periksa alasan, jumlah item, dan foto.", "Setujui atau tolak dengan catatan yang jelas.", "Ikuti proses refund/penggantian di luar sistem bila diperlukan."], [
                    ("Keputusan", "Approve / Reject", "Pilih sesuai hasil pemeriksaan."),
                    ("Catatan admin", "Disetujui setelah pemeriksaan", "Disarankan agar pelanggan dan audit memahami alasannya."),
                ], ["Permintaan terkait transaksi, pelanggan, dan item transaksi.", "Status retur masuk ke laporan retur/refund."], "31-retur")
    add_feature(doc, "Faktur Pajak", "Mengelola permintaan faktur pajak resmi yang diajukan pelanggan.",
                ["Filter permintaan yang berstatus requested.", "Buka detail dan verifikasi data wajib pajak.", "Tandai processing.", "Buat faktur pada sistem pajak resmi perusahaan.", "Unggah PDF, nomor, dan tanggal.", "Kirim email ke pelanggan bila siap."], [
                    ("Alasan penolakan", "NPWP tidak sesuai", "Wajib jika permintaan ditolak, maksimal 1.000 karakter."),
                    ("File faktur", "faktur-pajak.pdf", "Wajib saat upload, hanya PDF, maksimal 10 MB."),
                    ("Nomor faktur", "010.000-...", "Opsional, maksimal 100 karakter."),
                    ("Tanggal faktur", "2026-08-28", "Opsional, harus tanggal."),
                    ("Kirim email", "Ya", "Opsional saat upload; penerima mengikuti data permintaan."),
                ], ["Faktur pajak selalu terkait satu transaksi.", "Data NPWP adalah snapshot; perubahan profil pajak tidak mengubah permintaan lama.", "File hanya boleh diunduh oleh customer pemilik transaksi atau admin berizin."], "32-faktur-pajak",
                ["Invoice belanja biasa bukan faktur pajak resmi. Sistem ini menyimpan dan mengirim file yang sudah dibuat oleh tim finance."])
    add_feature(doc, "Product Reviews admin", "Memoderasi ulasan yang dikirim pelanggan.",
                ["Cari ulasan bermasalah.", "Sembunyikan/tampilkan sesuai pedoman.", "Hapus hanya jika memang diperlukan dan diizinkan."], [],
                ["Ulasan berasal dari item transaksi yang dibeli dan ditampilkan pada detail produk jika aktif."], "33-ulasan-produk")

    doc.add_heading("10. Panduan Admin — Penjualan B2B", level=1)
    add_note(doc, "Ringkasnya", "Quotation boleh dilewati. Jika harga sudah disepakati di luar sistem, admin dapat langsung membuat Sales Order. Setelah Sales Order terbentuk, alur fulfillment-nya sama.")
    add_feature(doc, "Quotations", "Dokumen penawaran harga untuk customer B2B. Harga dapat dinegosiasikan dan berlaku sampai tanggal tertentu.",
                ["Pilih customer existing atau isi customer manual.", "Tambahkan produk, qty, harga hasil negosiasi, dan catatan.", "Isi diskon, PPN, ongkir, biaya lain, dan masa berlaku.", "Simpan sebagai draft lalu kirim/ubah status sesuai respons customer.", "Jika accepted, convert sebagian atau seluruh qty menjadi Sales Order."], [
                    ("Customer", "PT Maju Jaya", "Wajib: pilih existing atau isi nama+telepon manual."),
                    ("Item", "Baut Hex M8", "Minimal satu varian."),
                    ("Qty", "100", "Wajib, bilangan bulat minimal 1."),
                    ("Harga nego", "750", "Wajib, 0 atau lebih."),
                    ("Diskon/biaya", "50000 / ongkir 100000", "Opsional, tidak boleh negatif."),
                    ("PPN", "11", "Opsional, 0–100."),
                    ("Berlaku hingga", "2026-09-15", "Wajib dan harus setelah hari ini."),
                    ("Catatan/syarat", "Harga franco Jakarta", "Opsional, maksimal 2.000 karakter."),
                ], ["Satu quotation dapat menghasilkan beberapa Sales Order melalui convert parsial.", "Jumlah yang sudah dikonversi dilacak per item.", "Quotation yang expired/rejected/closed menjadi hanya-baca dan tidak bisa dikonversi lagi."], "22-form-penawaran-b2b")
    add_screenshot(doc, "21-penawaran-b2b", "Tampilan daftar quotation")
    add_feature(doc, "Sales Orders", "Konfirmasi order B2B dan sumber utama proses pengiriman serta penagihan.",
                ["Buat dari quotation accepted atau langsung tanpa quotation.", "Untuk order langsung, isi customer, item, qty, harga, dan biaya.", "Setelah confirmed, buat Proforma bila perlu DP dan/atau buat Surat Jalan.", "Pantau jumlah terkirim dan sisa."], [
                    ("Customer", "PT Maju Jaya", "Wajib: existing atau manual."),
                    ("Item + qty + harga", "Baut M8 / 100 / 750", "Minimal satu item; qty minimal 1; harga minimal 0."),
                    ("PPN dan biaya", "11%, ongkir 100000", "Opsional, tidak boleh negatif."),
                    ("Catatan", "PO: MJ-001", "Opsional, maksimal 2.000 karakter."),
                ], ["SO dari quotation menyimpan referensi asal; SO langsung tidak punya quotation.", "Satu SO dapat memiliki banyak Surat Jalan, Proforma, dan Invoice.", "SO tidak ikut batal bila quotation asal kemudian ditutup/expired."], "24-form-sales-order")
    add_screenshot(doc, "23-sales-order", "Tampilan daftar Sales Order")
    add_feature(doc, "Proforma Invoices", "Dokumen tagihan awal/DP. Tidak mengatur pengiriman barang.",
                ["Buka Sales Order lalu buat Proforma.", "Pilih item dan jumlah yang ingin ditagihkan.", "Isi PPN dan biaya dokumen ini.", "Terbitkan dan catat pembayaran setiap kali uang masuk."], [
                    ("Item dan qty", "Baut M8 — 100", "Minimal satu item; isi 0 untuk baris yang tidak ditagih."),
                    ("PPN/ongkir/biaya", "11% / 0 / 0", "Opsional dan dihitung khusus untuk dokumen ini."),
                    ("Jumlah bayar", "5000000", "Wajib saat mencatat pembayaran, minimal 1."),
                    ("Tanggal bayar", "2026-08-28", "Wajib."),
                    ("Catatan", "DP transfer BCA", "Opsional, maksimal 500 karakter."),
                ], ["Pembayaran Proforma dicatat dalam ledger pembayaran.", "DP yang sudah dibayar menjadi kredit pengurang piutang Invoice B2B terkait.", "Status: draft, issued, partially_paid, paid, cancelled."], "25-proforma-invoice")
    add_feature(doc, "Surat Jalan dan Packing List", "Surat Jalan mencatat barang yang benar-benar dikirim. Packing List dibuat berpasangan untuk isi kemasan.",
                ["Buka Sales Order.", "Buat Surat Jalan.", "Isi penerima, alamat, kurir, jumlah paket, item, dan qty kirim.", "Pastikan stok cukup.", "Tandai shipped saat barang benar-benar keluar, lalu delivered ketika diterima."], [
                    ("Nama penerima", "Budi — Gudang PT Maju", "Wajib, maksimal 150 karakter."),
                    ("Alamat kirim", "Kawasan Industri...", "Wajib, maksimal 500 karakter."),
                    ("Kurir", "JNE Trucking", "Opsional, maksimal 100 karakter."),
                    ("Jumlah paket", "5", "Opsional, minimal 1."),
                    ("Item dan qty kirim", "Baut M8 — 40", "Minimal satu item; qty tidak boleh melebihi sisa SO."),
                ], ["Satu SO dapat dikirim bertahap lewat beberapa Surat Jalan.", "Saat shipped, stok dipotong dan stock movement dicatat.", "Packing List mengikuti item Surat Jalan dan merangkum berat/dimensi."], "26-surat-jalan",
                ["Surat Jalan yang sudah shipped tidak bisa dibatalkan lewat alur biasa. Koreksi setelah barang keluar harus mengikuti proses retur/kebijakan operasional."])
    add_screenshot(doc, "27-packing-list", "Tampilan daftar Packing List")
    add_feature(doc, "Invoice B2B", "Tagihan resmi B2B dengan jatuh tempo dan pencatatan pembayaran bertahap.",
                ["Buat dari Surat Jalan yang sudah dikirim/diterima, atau langsung dari item SO bila perlu tagih sebelum kirim.", "Pilih item/dokumen sumber.", "Isi tanggal jatuh tempo, PPN, dan biaya.", "Terbitkan lalu catat pembayaran sampai lunas."], [
                    ("Sumber invoice", "Surat Jalan / item Sales Order", "Wajib sesuai jalur yang dipilih."),
                    ("Jatuh tempo", "2026-09-27", "Wajib, hari ini atau sesudahnya."),
                    ("PPN", "11", "Opsional, 0–100."),
                    ("Biaya", "ongkir/admin/lain-lain", "Opsional, bilangan 0 atau lebih."),
                    ("Catatan biaya lain", "Biaya packing kayu", "Opsional, maksimal 255 karakter."),
                    ("Pembayaran", "10000000, 2026-09-01", "Jumlah minimal 1 dan tanggal wajib."),
                ], ["Invoice dapat menggabungkan Surat Jalan dalam satu Sales Order.", "Pembayaran tersimpan sebagai ledger dan mengubah status menjadi partially_paid/paid.", "Tanggal jatuh tempo dipakai untuk memantau piutang overdue."], "28-invoice-b2b")
    add_table(doc, ["Dokumen", "Status utama", "Aturan lanjut"], [
        ("Quotation", "draft, sent, accepted, partially_converted, rejected, expired, closed", "Hanya accepted/partially_converted yang belum kedaluwarsa dan masih punya sisa qty dapat di-convert."),
        ("Sales Order", "confirmed, partially_fulfilled, fulfilled, cancelled", "Bisa dibatalkan hanya selama belum ada Surat Jalan aktif."),
        ("Proforma", "draft, issued, partially_paid, paid, cancelled", "Pembayaran boleh bertahap."),
        ("Surat Jalan", "draft, shipped, delivered, cancelled", "Stok dipotong saat shipped; draft dapat dibatalkan."),
        ("Packing List", "mengikuti Surat Jalan", "Satu packing list untuk satu Surat Jalan."),
        ("Invoice B2B", "draft, issued, partially_paid, paid, cancelled", "Jatuh tempo wajib; pembayaran boleh bertahap."),
    ], [1.15, 2.35, 2.9])

    doc.add_heading("11. Reports", level=1)
    add_feature(doc, "Pusat laporan", "Kumpulan laporan untuk owner dan tim operasional.",
                ["Pilih laporan sesuai pertanyaan yang ingin dijawab.", "Atur periode dan filter company.", "Baca ringkasan, lalu ekspor bila tersedia dan diizinkan."], [],
                ["Semua laporan membaca transaksi/dokumen yang tersimpan; hasil tergantung status dan periode.", "Hak export dapat dipisahkan dari hak melihat laporan."], "34-pusat-laporan")
    add_feature(doc, "Jenis laporan", "Gunakan laporan yang paling cocok agar angka tidak disalahartikan.", [], [], [], "35-laporan-penjualan")
    add_table(doc, ["Laporan", "Kegunaan"], [
        ("Owner Overview", "Gambaran cepat penjualan dan kondisi utama bisnis."),
        ("Sales", "Nilai/order berdasarkan periode; dapat diekspor CSV bila punya izin."),
        ("Stock", "Stok tersedia, stok rendah, dan pergerakan."),
        ("Product Performance", "Produk terlaris atau performa penjualan produk."),
        ("Payment & Fulfillment", "Status pembayaran dan pemenuhan/pengiriman."),
        ("Customers", "Aktivitas dan nilai pelanggan."),
        ("Promo & Coupon", "Penggunaan promo/kupon."),
        ("Return & Refund", "Permintaan retur dan pengembalian dana."),
    ], [2.0, 4.4])

    doc.add_heading("12. Checklist Operasional", level=1)
    doc.add_heading("Sebelum toko dipakai", level=2)
    add_bullets(doc, [
        "Company aktif, nama toko, logo, dan prefix invoice sudah benar.",
        "Lokasi asal pengiriman sudah dipilih.",
        "Rekening transfer manual dan instruksinya sudah diuji.",
        "Tarif PPN sesuai kebijakan perusahaan.",
        "Kategori, produk, varian, harga, berat, stok, dan batas stok rendah sudah dicek.",
        "Role admin sudah mengikuti tugas; akses data sensitif tidak diberikan berlebihan.",
        "Email, Midtrans, RajaOngkir, WhatsApp, dan queue diuji oleh tim teknis/operasional.",
    ])
    doc.add_heading("Rutinitas harian", level=2)
    add_bullets(doc, [
        "Periksa order baru dan pembayaran yang perlu diverifikasi.",
        "Proses order sesuai antrean dan isi resi dengan benar.",
        "Pantau stok rendah dan pergerakan stok yang tidak biasa.",
        "Periksa permintaan retur dan faktur pajak.",
        "Pantau piutang Invoice B2B yang mendekati/lewat jatuh tempo.",
        "Pastikan company aktif sebelum membuat dokumen baru.",
    ])
    doc.add_heading("Sebelum menghapus/menonaktifkan data", level=2)
    add_bullets(doc, [
        "Cek apakah data sudah dipakai produk, transaksi, promo, atau dokumen B2B.",
        "Untuk data berhistori, lebih aman ubah menjadi inactive daripada menghapus.",
        "Simpan alasan koreksi stok, pembatalan, penolakan retur, dan penolakan faktur pajak.",
    ])

    doc.add_heading("13. Masalah yang Sering Terjadi", level=1)
    add_table(doc, ["Masalah", "Yang perlu dicek"], [
        ("Produk tidak muncul", "Status produk aktif, company aktif, kategori, varian, dan stok."),
        ("Ongkir tidak muncul", "Lokasi toko, provinsi/kota tujuan, berat varian, dan koneksi RajaOngkir."),
        ("Kupon ditolak", "Periode, status, minimum belanja, kuota, member-only, dan company."),
        ("Tidak bisa convert quotation", "Status harus accepted/partially_converted, belum expired/closed, dan masih ada sisa qty."),
        ("Tidak bisa buat Surat Jalan", "SO aktif, qty tidak melebihi sisa, dan varian masih tersedia."),
        ("Pembayaran belum mengubah status", "Jumlah/tanggal, dokumen yang dipilih, dan apakah saldo sudah benar-benar lunas."),
        ("Menu admin hilang", "Role/permission dan penugasan company pengguna admin."),
        ("Gambar gagal diunggah", "Format dan ukuran file sesuai batas pada tabel panduan."),
        ("Lacak pesanan terkunci", "Tunggu hingga batas 15 menit selesai; pastikan email dan nomor order benar."),
    ], [2.1, 4.3])

    doc.add_heading("14. Lampiran — Daftar Fitur yang Dipetakan", level=1)
    modules = [
        ("Pelanggan", "Beranda, kategori, pencarian, detail produk, flash sale, cart, checkout, pembayaran, guest tracking, login/register/reset, profil, alamat, pesanan, wishlist, notifikasi, poin, retur, ulasan, faktur pajak, blog, promo page, content page, newsletter, API katalog"),
        ("Admin utama", "Dashboard, company switch, customers, membership tiers, products/import, categories, variants, stock"),
        ("Transaksi B2C", "Buat transaksi manual, daftar/detail transaksi, verifikasi pembayaran, proses/kirim/selesai/batal, label, retur, faktur pajak, ulasan"),
        ("Promo & konten", "Flash sale, coupons, banners, newsletter, promo pages, content website"),
        ("B2B", "Quotations, Sales Orders langsung/convert, Proforma Invoices, Surat Jalan, Packing Lists, Invoice B2B, pembayaran bertahap"),
        ("Laporan", "Owner, sales, stock, product performance, payments/fulfillment, customers, promos/coupons, returns/refunds"),
        ("Pengaturan", "Store identity, location, manual payment, tax, WhatsApp gateway, social media, API docs, admin users, roles & permissions, change password"),
        ("Integrasi", "Midtrans, RajaOngkir, mail, WhatsApp gateway, Google login, public catalog API"),
    ]
    add_table(doc, ["Kelompok", "Cakupan"], modules, [1.35, 5.05])
    add_note(doc, "Ruang lingkup", "Dokumen ini memetakan fitur yang tersedia pada source code dan menu aktif per 28 Agustus 2026. Menu contoh/template yang sengaja dikomentari di konfigurasi tidak diperlakukan sebagai fitur operasional.")

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build_manual()
    print(output)
